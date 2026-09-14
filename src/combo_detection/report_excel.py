"""Builds the combination-ruleset detection report as an Excel workbook --
same severity-color conventions as SOD_Detection's own reports (F8DEDE/
FBE5DB/FDF1D9 for High/Medium/Low), a dark bold header row, frozen panes,
and autofilter.

Three sheets:
  Read Me            - what this is, DRAFT status, summary counts
  Findings            - one row per (user, combo) match, wide format
  Manual Controls      - rules that can never be detected via T-code at all
"""

import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1B2430")
HEADER_FONT = Font(color="FFFFFF", bold=True)
LEVEL_FILL = {"High": PatternFill("solid", fgColor="F8DEDE"), "Medium": PatternFill("solid", fgColor="FBE5DB"), "Low": PatternFill("solid", fgColor="FDF1D9")}


def _finding_columns(max_functions):
    """Build the Findings-sheet column spec for a ruleset whose widest rule
    has `max_functions` functions -- computed per-run (see build_excel_report)
    rather than a fixed constant, so a rule with more legs than any seen so
    far still gets a column instead of being silently dropped.
    """
    columns = [
        ("combo_id", "Combo ID", 14),
        ("process_area", "Module", 12),
        ("category", "Category", 22),
        ("risk_level", "Risk", 9),
        ("control_type", "Control Type", 12),
        ("user_id", "User", 12),
    ]
    for i in range(1, max_functions + 1):
        columns += [
            (f"function_{i}_name", f"Function {i}", 20),
            (f"function_{i}_tcodes", f"F{i} T-code(s)", 12),
            (f"function_{i}_roles", f"F{i} Granting Role(s)", 26),
        ]
    columns += [
        ("business_impact", "Business Impact", 42),
        ("suggested_mitigation", "Suggested Mitigation", 42),
        ("status", "Ruleset Status", 30),
    ]
    return columns

MANUAL_COLUMNS = [
    ("combo_id", "Combo ID", 14),
    ("process_area", "Module", 12),
    ("category", "Category", 22),
    ("risk_level", "Risk", 9),
    ("functions_with_no_tcode", "Function(s) with no T-code", 40),
    ("business_impact", "Business Impact", 42),
    ("suggested_mitigation", "Suggested Mitigation", 42),
]


def _write_header(ws, columns):
    for col_idx, (_, header, width) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"


def _finding_row_values(finding, max_functions):
    values = {
        "combo_id": finding["combo_id"],
        "process_area": finding["process_area"],
        "category": finding["category"],
        "risk_level": finding["risk_level"],
        "control_type": finding["control_type"],
        "user_id": finding["user_id"],
        "business_impact": finding["business_impact"],
        "suggested_mitigation": finding["suggested_mitigation"],
        "status": finding["status"],
    }
    for i in range(max_functions):
        if i < len(finding["functions"]):
            fn = finding["functions"][i]
            values[f"function_{i + 1}_name"] = fn["name"]
            values[f"function_{i + 1}_tcodes"] = "|".join(fn["matched_tcodes"])
            values[f"function_{i + 1}_roles"] = ", ".join(fn["granting_roles"]) or "(unknown)"
        else:
            values[f"function_{i + 1}_name"] = ""
            values[f"function_{i + 1}_tcodes"] = ""
            values[f"function_{i + 1}_roles"] = ""
    return values


def build_excel_report(rules, findings, output_path, generated_for=""):
    wb = Workbook()

    # --- Read Me ---
    cover = wb.active
    cover.title = "Read Me"
    cover.sheet_view.showGridLines = False
    cover.column_dimensions["A"].width = 105

    checkable = [r for r in rules if all(fn["tcodes"] for fn in r["functions"])]
    affected_users = sorted({f["user_id"] for f in findings})
    lines = [
        ("SOD Combination Ruleset -- Detection Report", 16, True),
        ("", None, False),
        ("DRAFT ruleset -- not approved by IT Governance. These combination rules are proposed", 11, True),
        ("escalations of the existing 144-rule SOD ruleset, not transcribed policy. Any finding on", 11, False),
        ("the Findings sheet is a candidate for review, not a confirmed violation.", 11, False),
        ("", None, False),
        (f"Generated: {datetime.date.today():%d %B %Y}" + (f"   |   For: {generated_for}" if generated_for else ""), 11, False),
        (f"Rules in ruleset: {len(rules)}   |   Checkable via T-code data: {len(checkable)}", 11, False),
        (f"Findings: {len(findings)}   |   Users affected: {len(affected_users)}", 11, True),
        ("", None, False),
        ("Findings sheet: one row per (user, rule) match -- every function in the rule matched", 11, False),
        ("that user's access, not just some of them.", 11, False),
        ("Manual Controls sheet: rules with at least one function that has no T-code mapping at", 11, False),
        ("all -- these can never produce a finding here and need organizational tracking instead", 11, False),
        ("(same principle as SOD_Detection's own manual-controls checklist).", 11, False),
    ]
    for i, (text, size, bold) in enumerate(lines, start=1):
        cell = cover.cell(row=i, column=1, value=text)
        cell.font = Font(size=size or 11, bold=bold)
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    # --- Findings ---
    # Sized off the ruleset's widest rule, not findings -- a rule with more
    # legs than any match found yet must still get a column when it does.
    max_functions = max((len(r["functions"]) for r in rules), default=0)
    finding_columns = _finding_columns(max_functions)

    ws = wb.create_sheet("Findings")
    _write_header(ws, finding_columns)
    for row_idx, finding in enumerate(findings, start=2):
        values = _finding_row_values(finding, max_functions)
        fill = LEVEL_FILL.get(finding["risk_level"])
        for col_idx, (key, _, _) in enumerate(finding_columns, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=values.get(key, ""))
            if fill:
                cell.fill = fill
            cell.alignment = Alignment(vertical="top", wrap_text=key in ("business_impact", "suggested_mitigation", "status"))
    if findings:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(finding_columns))}{len(findings) + 1}"

    # --- Manual Controls ---
    ws2 = wb.create_sheet("Manual Controls")
    _write_header(ws2, MANUAL_COLUMNS)
    manual_rules = [r for r in rules if any(not fn["tcodes"] for fn in r["functions"])]
    for row_idx, rule in enumerate(manual_rules, start=2):
        no_tcode_fns = [fn["name"] for fn in rule["functions"] if not fn["tcodes"]]
        values = {
            "combo_id": rule["combo_id"],
            "process_area": rule["process_area"],
            "category": rule["category"],
            "risk_level": rule["risk_level"],
            "functions_with_no_tcode": ", ".join(no_tcode_fns),
            "business_impact": rule["business_impact"],
            "suggested_mitigation": rule["suggested_mitigation"],
        }
        fill = LEVEL_FILL.get(rule["risk_level"])
        for col_idx, (key, _, _) in enumerate(MANUAL_COLUMNS, start=1):
            cell = ws2.cell(row=row_idx, column=col_idx, value=values.get(key, ""))
            if fill:
                cell.fill = fill
            cell.alignment = Alignment(vertical="top", wrap_text=key in ("business_impact", "suggested_mitigation"))
    if manual_rules:
        ws2.auto_filter.ref = f"A1:{get_column_letter(len(MANUAL_COLUMNS))}{len(manual_rules) + 1}"

    wb.save(output_path)
