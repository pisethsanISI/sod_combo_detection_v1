import openpyxl

from combo_detection.access import build_user_tcode_roles
from combo_detection.detect import detect_combination_conflicts
from combo_detection.report_excel import build_excel_report
from combo_detection.ruleset import load_combination_ruleset, rules_needing_manual_tracking


def test_build_excel_report_has_expected_sheets_and_findings(tmp_path):
    rules = load_combination_ruleset("ruleset/sap_sod_combination_ruleset.json")
    user_tcode_roles = build_user_tcode_roles(
        "data/sample_role_tcodes.csv", "data/sample_user_roles.csv", "data/sample_composite_roles.csv"
    )
    user_tcodes = {u: set(t.keys()) for u, t in user_tcode_roles.items()}
    findings = detect_combination_conflicts(rules, user_tcodes, user_tcode_roles)

    out_path = str(tmp_path / "report.xlsx")
    build_excel_report(rules, findings, out_path, generated_for="Test Run")

    wb = openpyxl.load_workbook(out_path)
    assert wb.sheetnames == ["Read Me", "Findings", "Manual Controls"]

    findings_ws = wb["Findings"]
    assert findings_ws.max_row == len(findings) + 1  # header + rows
    header = [c.value for c in findings_ws[1]]
    assert "Combo ID" in header
    assert "Function 1" in header
    assert "F1 Granting Role(s)" in header

    # the planted U1 finding (COMBO-P2P-01) should be present
    combo_ids = [row[0].value for row in findings_ws.iter_rows(min_row=2)]
    assert "COMBO-P2P-01" in combo_ids


def test_manual_controls_sheet_lists_rules_with_a_blank_function(tmp_path):
    rules = load_combination_ruleset("ruleset/sap_sod_combination_ruleset.json")
    expected_manual = {r["combo_id"] for r in rules_needing_manual_tracking(rules)}

    out_path = str(tmp_path / "report.xlsx")
    build_excel_report(rules, [], out_path)

    wb = openpyxl.load_workbook(out_path)
    ws = wb["Manual Controls"]
    listed = {row[0].value for row in ws.iter_rows(min_row=2)}

    assert listed == expected_manual
    assert "COMBO-TR-02" in listed


def test_build_excel_report_handles_zero_findings(tmp_path):
    rules = load_combination_ruleset("ruleset/sap_sod_combination_ruleset.json")
    out_path = str(tmp_path / "empty.xlsx")

    build_excel_report(rules, [], out_path)

    wb = openpyxl.load_workbook(out_path)
    assert wb["Findings"].max_row == 1  # header only


def test_build_excel_report_does_not_truncate_a_rule_with_more_than_three_functions(tmp_path):
    # Regression test: report_excel.py used to size the Findings sheet off a
    # hardcoded MAX_FUNCTIONS = 3, silently dropping any 4th+ function's
    # columns even though detect.py's matching is fully N-way generic.
    rules = [
        {
            "combo_id": "COMBO-TEST-4WAY",
            "process_area": "P2P",
            "category": "Test",
            "risk_level": "High",
            "control_type": "System",
            "functions": [
                {"name": "A", "tcodes": {"T1"}},
                {"name": "B", "tcodes": {"T2"}},
                {"name": "C", "tcodes": {"T3"}},
                {"name": "D", "tcodes": {"T4"}},
            ],
            "business_impact": "x",
            "suggested_mitigation": "y",
            "status": "DRAFT",
            "notes": "",
        }
    ]
    findings = [
        {
            "user_id": "U1",
            "combo_id": "COMBO-TEST-4WAY",
            "process_area": "P2P",
            "category": "Test",
            "risk_level": "High",
            "control_type": "System",
            "functions": [
                {"name": "A", "matched_tcodes": ["T1"], "granting_roles": ["R1"]},
                {"name": "B", "matched_tcodes": ["T2"], "granting_roles": ["R2"]},
                {"name": "C", "matched_tcodes": ["T3"], "granting_roles": ["R3"]},
                {"name": "D", "matched_tcodes": ["T4"], "granting_roles": ["R4"]},
            ],
            "business_impact": "x",
            "suggested_mitigation": "y",
            "status": "DRAFT",
            "notes": "",
        }
    ]

    out_path = str(tmp_path / "four_way.xlsx")
    build_excel_report(rules, findings, out_path)

    wb = openpyxl.load_workbook(out_path)
    ws = wb["Findings"]
    header = [c.value for c in ws[1]]
    assert "Function 4" in header
    assert "F4 Granting Role(s)" in header

    row = dict(zip(header, [c.value for c in ws[2]]))
    assert row["Function 4"] == "D"
    assert row["F4 T-code(s)"] == "T4"
    assert row["F4 Granting Role(s)"] == "R4"
